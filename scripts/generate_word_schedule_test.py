from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from schedule_test_data import CourseBlock, ScheduleVariant, get_variant


def set_run_font(run, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), "D9D9D9")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def remove_paragraph_borders(paragraph) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def write_cell(cell, text: str, *, centered: bool = False, white: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1

    for index, line in enumerate(text.splitlines()):
        if index > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, 8.5, bold=index == 0, color="FFFFFF" if white else "000000")

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)


def format_course_offering(course: CourseBlock) -> str:
    week_text = (
        f"第{course.first_week}周"
        if course.first_week == course.last_week
        else f"{course.first_week}-{course.last_week}周"
    )
    if course.parity == "Odd":
        week_text += "(单)"
    elif course.parity == "Even":
        week_text += "(双)"
    period_text = (
        f"第{course.first_period}节"
        if course.first_period == course.last_period
        else f"第{course.first_period}-{course.last_period}节"
    )
    return f"{week_text}[讲授] {course.teacher} {course.location} {period_text}"


def build_document(variant: ScheduleVariant) -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title_style = document.styles["Title"]
    title_properties = title_style._element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run(f"课程表导入测试样本 {variant.key}")
    set_run_font(title_run, 18, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    remove_paragraph_borders(title)

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    intro.paragraph_format.line_spacing = 1.15
    run = intro.add_run(
        f"本文件用于验证课程表导入。每节固定 {variant.lesson_minutes} 分钟，"
        f"首节从 {variant.periods[0].start} 开始，最后一节在 {variant.periods[-1].end} 结束。"
    )
    set_run_font(run, 9)

    table = document.add_table(rows=1 + len(variant.periods), cols=6)
    table.autofit = False
    table.allow_autofit = False

    header_labels = ["节次 / 星期", "星期一", "星期二", "星期三", "星期四", "星期五"]
    for column, label in enumerate(header_labels):
        write_cell(table.cell(0, column), label, centered=True, white=True)
        shade_cell(table.cell(0, column), "1F4E79")
    set_repeat_header(table.rows[0])

    for period_index, period in enumerate(variant.periods, start=1):
        write_cell(
            table.cell(period_index, 0),
            f"第{period.number}节 ({period.start}-{period.end})",
            centered=True,
        )
        shade_cell(table.cell(period_index, 0), "EAF1F8" if period_index % 2 else "FFFFFF")
        for column in range(1, 6):
            write_cell(table.cell(period_index, column), "")
            shade_cell(table.cell(period_index, column), "F7FAFD" if period_index % 2 else "FFFFFF")

    for course in variant.courses:
        merged = table.cell(course.first_period, course.weekday + 1).merge(
            table.cell(course.last_period, course.weekday + 1)
        )
        write_cell(
            merged,
            f"{course.name}[{course.code}] {course.credits:g}学分\n"
            f"{format_course_offering(course)}",
        )
        shade_cell(merged, "EAF1F8" if course.first_period % 2 else "FFFFFF")

    widths = [1.05, 1.9, 1.9, 1.9, 1.9, 1.9]
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Inches(0.28)
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)

    return document


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: generate_word_schedule_test.py <variant> <output.docx>")

    variant = get_variant(sys.argv[1])
    output_path = Path(sys.argv[2]).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    build_document(variant).save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
